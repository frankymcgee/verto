"""Read the INX Summary Events export without treating response notes as controls."""

import hashlib
import json
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile


SOURCE_SYSTEM = "INX InControl"
MAX_EVENTS = 2000
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
SHEET_NAME = "v_EventsReport"
# Personnel columns are deliberately excluded, including from fingerprints.
PERSONNEL_COLUMNS = {
    "reported by name", "originator", "responsible manager(s)",
    "closed out by", "reported by lookup name",
}
COLUMN_FIELDS = {
    "reference": "source_reference",
    "event date": "event_datetime",
    "event type": "event_type",
    "event sub type": "event_sub_type",
    "re workgroup name": "re_workgroup",
    "location": "location",
    "country": "country",
    "immediate action taken": "immediate_actions",
    "status": "event_status",
    "short observation": "short_observation",
    "detailed observation": "detailed_observation",
    "moderator comment": "moderator_comment",
    "review date": "review_date",
    "reviewed date": "reviewed_date",
    "review summary": "review_summary",
    "closed out date": "closed_out_datetime",
    "workgroup": "workgroup",
}
REQUIRED_COLUMNS = {"reference", "event date", "short observation", "detailed observation", "immediate action taken"}
DATE_FIELDS = {"event_datetime", "review_date", "reviewed_date", "closed_out_datetime"}
LONG_FIELDS = {"immediate_actions", "short_observation", "detailed_observation", "moderator_comment", "review_summary"}


def source_key(system, reference):
    system, reference = str(system or "").strip(), str(reference or "").strip()
    if not system or not reference:
        raise ValueError("Source system and incident reference are required.")
    return hashlib.sha256(f"{system.casefold()}\0{reference.casefold()}".encode()).hexdigest()


def source_fingerprint(values):
    payload = {field: str(values.get(field) or "").strip() for field in COLUMN_FIELDS.values()}
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _datetime(value, field, row):
    if value is None or value == "":
        if field == "event_datetime":
            raise ValueError(f"Row {row}: Event Date is required.")
        return None
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, datetime.min.time())
    elif isinstance(value, str):
        result = None
        # INX text dates use Australian day/month order. Excel dates are already typed.
        for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S",
                    "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
            try:
                result = datetime.strptime(value.strip(), fmt)
                break
            except ValueError:
                continue
        if result is None:
            raise ValueError(f"Row {row}: invalid date in {field}.")
    else:
        raise ValueError(f"Row {row}: invalid date in {field}.")
    if result.tzinfo is not None:
        raise ValueError(f"Row {row}: use the original local INX date and time without a timezone conversion.")
    return result.isoformat(sep=" ")


def parse_rows(rows):
    """Validate the complete export before any database writes; return restricted source data."""
    iterator = iter(rows)
    header = next(iterator, ())
    columns = [str(value or "").strip().lstrip("\ufeff").casefold() for value in header]
    if len([c for c in columns if c]) != len(set(c for c in columns if c)):
        raise ValueError("The export has duplicate column headings.")
    missing = REQUIRED_COLUMNS - set(columns)
    if missing:
        raise ValueError("Missing INX columns: " + ", ".join(sorted(missing)) + ".")
    records, identities = [], set()
    for row_number, cells in enumerate(iterator, start=2):
        if not any(value is not None and value != "" for value in cells):
            continue
        if len(records) >= MAX_EVENTS:
            raise ValueError(f"Import at most {MAX_EVENTS} events at a time.")
        source = dict(zip(columns, cells))
        record = {}
        for column, field in COLUMN_FIELDS.items():
            value = source.get(column)
            if field in DATE_FIELDS:
                record[field] = _datetime(value, field, row_number)
            else:
                if isinstance(value, (bool, date, datetime)):
                    raise ValueError(f"Row {row_number}: invalid text in {column}.")
                if field == "source_reference" and isinstance(value, float) and value.is_integer():
                    value = int(value)
                record[field] = str(value if value is not None else "").strip()
                limit = 32767 if field in LONG_FIELDS else 140
                if len(record[field]) > limit:
                    raise ValueError(f"Row {row_number}: {column} exceeds {limit} characters.")
        if not record["source_reference"]:
            raise ValueError(f"Row {row_number}: Reference is required.")
        key = source_key(SOURCE_SYSTEM, record["source_reference"])
        if key in identities:
            raise ValueError(f"Row {row_number}: duplicate incident reference.")
        identities.add(key)
        record["source_key"] = key
        record["source_fingerprint"] = source_fingerprint(record)
        records.append(record)
    if not records:
        raise ValueError("The INX export contains no events.")
    return {
        "records": records,
        "ignored_columns": [str(header[i]) for i, column in enumerate(columns)
                            if column and column not in COLUMN_FIELDS],
        "missing_review_summaries": sum(not record["review_summary"] for record in records),
    }


def _open_workbook(content):
    if not isinstance(content, bytes) or len(content) > MAX_FILE_BYTES:
        raise ValueError("Upload an XLSX export of no more than 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("The expanded workbook is too large.")
        # openpyxl is supplied by Frappe v16; imported lazily for row-only validation.
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("The file is not a readable INX XLSX workbook.") from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"The INX export must contain the {SHEET_NAME} worksheet.")
        sheet = workbook[SHEET_NAME]
        if sheet.max_column > 64 or sheet.max_row > MAX_EVENTS + 1:
            raise ValueError(f"Import at most {MAX_EVENTS} events and 64 columns at a time.")
        return workbook
    except Exception:
        workbook.close()
        raise


def parse_workbook(content):
    """Read XLSX only. Never evaluate formulas, external links or macros."""
    workbook = _open_workbook(content)
    try:
        sheet = workbook[SHEET_NAME]

        def rows():
            for cells in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("The INX export must contain values, not formulas.")
                yield [cell.value for cell in cells]

        return parse_rows(rows())
    finally:
        workbook.close()


def personnel_hints(content, reference):
    """Read this incident's excluded name columns in memory only for redaction."""
    workbook = _open_workbook(content)
    try:
        rows = iter(workbook[SHEET_NAME].iter_rows())
        columns = [str(cell.value or "").strip().lstrip("\ufeff").casefold() for cell in next(rows, ())]
        if "reference" not in columns or len(columns) != len(set(columns)):
            raise ValueError("Invalid INX column headings.")
        ref_column = columns.index("reference")
        name_columns = [i for i, column in enumerate(columns) if column in PERSONNEL_COLUMNS]
        for cells in rows:
            value = cells[ref_column].value
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            if str(value or "").strip().casefold() != str(reference).strip().casefold():
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("Personnel hints must contain values, not formulas.")
            return [str(cells[i].value).strip() for i in name_columns if cells[i].value]
        return []
    finally:
        workbook.close()
