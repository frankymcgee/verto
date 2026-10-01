import hashlib
from datetime import datetime
from io import BytesIO
from unittest import TestCase
from unittest.mock import Mock, patch

import frappe
from frappe.tests import IntegrationTestCase

from verto.api import inx_incident_import as importer
from verto.api.mobile import voice_jha_incidents
from verto.safety.inx_import import SHEET_NAME, parse_rows, source_fingerprint


def synthetic_export_rows(reference):
    return [
        ['Reference', 'Event Date', 'Short Observation', 'Detailed Observation', 'Immediate Action Taken', 'Status', 'Closed Out Date'],
        [reference, '06/01/2025 03:30', 'Hand caught between flanges', 'Hand entered a pinch point.', 'Stopped the task.', 'Closed', datetime(2025, 1, 7, 10, 5, 7, 123000)],
    ]


def synthetic_export(reference='TEST-INX-101'):
    return parse_rows(synthetic_export_rows(reference))


class TestINXImportEndpoint(TestCase):
    def setUp(self):
        self.manager_check = importer._require_manager
        self.manager = self.enterContext(patch.object(importer, '_require_manager'))
        self.file = Mock(is_private=1, file_url='/private/files/test-inx.xlsx', file_name='test-inx.xlsx')
        self.file.get_content.return_value = b'synthetic export'
        self.get_doc = self.enterContext(patch.object(frappe, 'get_doc', return_value=self.file))
        self.parse = self.enterContext(patch.object(importer, 'parse_workbook', return_value=synthetic_export()))
        self.existing = self.enterContext(patch.object(importer, '_existing', return_value={}))
        self.upsert = self.enterContext(patch.object(importer, '_upsert'))

    def test_preview_has_counts_only_and_makes_no_writes(self):
        result = importer.import_inx_export('FILE-TEST')
        self.assertEqual((result['row_count'], result['new_sources'], result['new_lessons']), (1, 1, 1))
        self.assertEqual(result['missing_review_summaries'], 1)
        self.assertFalse(result['investigation_actions_provided'])
        self.assertNotIn('records', result)
        self.file.check_permission.assert_called_once_with('read')
        self.upsert.assert_not_called()
        self.file.save.assert_not_called()

    def test_public_files_and_missing_preview_are_rejected(self):
        self.file.is_private = 0
        with self.assertRaises(frappe.ValidationError):
            importer.import_inx_export('FILE-TEST')
        self.file.is_private = 1
        with self.assertRaises(frappe.ValidationError):
            importer.import_inx_export('FILE-TEST', dry_run=False)
        self.parse.assert_not_called()
        self.upsert.assert_not_called()

    def test_failed_import_rolls_back_the_entire_batch(self):
        self.upsert.side_effect = RuntimeError('Synthetic failure')
        with patch.object(frappe.db, 'savepoint'), patch.object(frappe.db, 'rollback') as rollback:
            with self.assertRaises(RuntimeError):
                importer.import_inx_export('FILE-TEST', dry_run=False, expected_sha256=hashlib.sha256(b'synthetic export').hexdigest())
        rollback.assert_called_once_with(save_point='inx_incident_import')
        self.file.save.assert_not_called()

    def test_crew_role_cannot_import_or_read_the_file(self):
        with patch.object(importer, '_require_manager', self.manager_check), patch.object(frappe, 'get_roles', return_value=['All']), patch.object(frappe, 'session', frappe._dict(user='crew@example.test')):
            with self.assertRaises(frappe.PermissionError):
                importer.import_inx_export('FILE-TEST')
        self.get_doc.assert_not_called()


class IntegrationTestINXIncidentImport(IntegrationTestCase):
    def _upload_workbook(self):
        from openpyxl import Workbook

        references = [f'TEST-INX-{frappe.generate_hash(length=10)}' for _ in range(2)]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(synthetic_export_rows(references[0])[0])
        for reference in references:
            sheet.append(synthetic_export_rows(reference)[1])
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        uploaded = frappe.get_doc({
            'doctype': 'File', 'file_name': f'{references[0]}.xlsx',
            'is_private': 1, 'content': content.getvalue(),
        }).insert()
        keys = [synthetic_export(reference)['records'][0]['source_key'] for reference in references]
        return uploaded, keys

    def test_private_workbook_import_handles_automatic_attachment_updates(self):
        uploaded, keys = self._upload_workbook()
        preview = importer.import_inx_export(uploaded.name)
        self.assertEqual((preview['row_count'], preview['new_sources'], preview['new_lessons']), (2, 2, 2))
        self.assertFalse(frappe.db.get_value('File', uploaded.name, 'attached_to_name'))

        result = importer.import_inx_export(uploaded.name, dry_run=False, expected_sha256=preview['file_sha256'])
        self.assertFalse(result['dry_run'])
        uploaded.reload()
        self.assertEqual(uploaded.is_private, 1)
        self.assertEqual(uploaded.attached_to_doctype, importer.SOURCE_DOCTYPE)
        self.assertEqual(uploaded.attached_to_name, keys[0])
        self.assertEqual(uploaded.attached_to_field, 'source_file')
        for key in keys:
            source = frappe.get_doc(importer.SOURCE_DOCTYPE, key)
            lesson = frappe.get_doc(importer.LEARNING_DOCTYPE, key)
            self.assertEqual(source.source_file, uploaded.file_url)
            self.assertEqual(lesson.available_for_jha, 0)
            self.assertEqual(lesson.source_review_required, 1)

        attachments = frappe.get_all('File', filters={'file_url': uploaded.file_url},
                                     fields=['is_private', 'attached_to_doctype', 'attached_to_name'])
        self.assertTrue(attachments)
        for attachment in attachments:
            self.assertEqual(attachment.is_private, 1)
            self.assertEqual(attachment.attached_to_doctype, importer.SOURCE_DOCTYPE)
            self.assertIn(attachment.attached_to_name, keys)

        repeated = importer.import_inx_export(uploaded.name, dry_run=False, expected_sha256=preview['file_sha256'])
        self.assertEqual((repeated['new_sources'], repeated['new_lessons'], repeated['unchanged_sources']), (0, 0, 2))

    def test_final_file_save_failure_rolls_back_attachments_and_allows_retry(self):
        uploaded, keys = self._upload_workbook()
        preview = importer.import_inx_export(uploaded.name)
        with patch('frappe.core.doctype.file.file.File.save', side_effect=RuntimeError('Synthetic attachment failure')):
            with self.assertRaisesRegex(RuntimeError, 'Synthetic attachment failure'):
                importer.import_inx_export(uploaded.name, dry_run=False, expected_sha256=preview['file_sha256'])

        for key in keys:
            self.assertFalse(frappe.db.exists(importer.SOURCE_DOCTYPE, key))
            self.assertFalse(frappe.db.exists(importer.LEARNING_DOCTYPE, key))
        uploaded.reload()
        self.assertEqual(uploaded.is_private, 1)
        self.assertFalse(uploaded.attached_to_doctype)
        self.assertFalse(uploaded.attached_to_name)
        self.assertFalse(uploaded.attached_to_field)
        self.assertEqual(frappe.db.count('File', {'file_url': uploaded.file_url}), 1)

        retried = importer.import_inx_export(uploaded.name, dry_run=False, expected_sha256=preview['file_sha256'])
        self.assertEqual((retried['new_sources'], retried['new_lessons']), (2, 2))

    def test_repeat_import_preserves_curation_and_changed_source_withdraws_it(self):
        reference = f'TEST-INX-{frappe.generate_hash(length=10)}'
        record = synthetic_export(reference)['records'][0]
        key = record['source_key']
        importer._upsert(record, None)
        source = frappe.get_doc(importer.SOURCE_DOCTYPE, key)
        lesson = frappe.get_doc(importer.LEARNING_DOCTYPE, key)
        self.assertEqual(lesson.available_for_jha, 0)
        self.assertEqual(lesson.source_review_required, 1)
        self.assertFalse(lesson.incident_summary)
        self.assertFalse(lesson.investigation_actions)
        self.assertFalse(lesson.immediate_actions)
        self.assertEqual(lesson.inx_source, source.name)
        self.assertEqual(source.source_fingerprint, record['source_fingerprint'])
        self.assertIn('Pinch Points', source.suggested_mechanisms)
        lesson.title = 'Flange alignment'
        lesson.incident_summary = 'Hands entered the pinch point during alignment.'
        lesson.immediate_actions = 'Stopped the task.'
        lesson.recommended_controls = 'Reviewed alignment tooling.'
        lesson.available_for_jha = 1
        lesson.save()
        self.assertEqual(lesson.source_review_required, 0)
        approved_modified = str(lesson.modified)
        source.reload()
        self.assertEqual(str(source.closed_out_datetime), '2025-01-07 10:05:07.123000')
        source.save()
        importer._upsert(record, None)
        lesson.reload()
        self.assertEqual(lesson.available_for_jha, 1)
        self.assertEqual(str(lesson.modified), approved_modified)
        self.assertEqual(lesson.recommended_controls, 'Reviewed alignment tooling.')
        changed = dict(record, detailed_observation='Unexpected movement trapped a hand.')
        changed['source_fingerprint'] = source_fingerprint(changed)
        importer._upsert(changed, None)
        lesson.reload()
        self.assertEqual(lesson.available_for_jha, 0)
        self.assertEqual(lesson.source_review_required, 1)
        self.assertEqual(lesson.incident_summary, 'Hands entered the pinch point during alignment.')
        self.assertEqual(lesson.recommended_controls, 'Reviewed alignment tooling.')
        results = voice_jha_incidents.find_incident_learning('Pinch point')
        self.assertNotIn(key, [row['name'] for row in results['incidents']])

    def test_changed_source_rejects_stale_approval_and_raw_fields_are_manager_only(self):
        reference = f'TEST-INX-{frappe.generate_hash(length=10)}'
        record = synthetic_export(reference)['records'][0]
        importer._upsert(record, None)
        lesson = frappe.get_doc(importer.LEARNING_DOCTYPE, record['source_key'])
        source = frappe.get_doc(importer.SOURCE_DOCTYPE, record['source_key'])
        source.detailed_observation = 'Unexpected movement at the flange.'
        source.save()
        # Simulate a stale fingerprint even on a freshly loaded lesson.
        lesson.reload()
        lesson.inx_source_fingerprint = record['source_fingerprint']
        lesson.incident_summary = 'Reviewed crew summary.'
        lesson.available_for_jha = 1
        with self.assertRaises(frappe.ValidationError):
            lesson.save()
        self.assertNotIn('All', {permission.role for permission in frappe.get_meta(importer.SOURCE_DOCTYPE).permissions})
        original_user = frappe.session.user
        try:
            frappe.set_user('Guest')
            self.assertFalse(frappe.has_permission(importer.SOURCE_DOCTYPE, 'read'))
            lesson.apply_fieldlevel_read_permissions()
            self.assertFalse(lesson.get('inx_source'))
        finally:
            frappe.set_user(original_user)
        self.assertNotIn('detailed_observation', voice_jha_incidents.FIELDS)
